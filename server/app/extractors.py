import csv
import mimetypes
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".jsonl",
    ".csv",
    ".tsv",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".css",
    ".sql",
}


def guess_mime(path: Path) -> str:
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def extract_text(path: Path, max_chars: int = 80_000) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_EXTENSIONS:
            return path.read_text(encoding="utf-8", errors="replace")[:max_chars]
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join((page.extract_text() or "") for page in reader.pages[:100])[:max_chars]
        if suffix == ".docx":
            document = Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)[:max_chars]
        if suffix in {".xlsx", ".xlsm"}:
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            chunks: list[str] = []
            for sheet in workbook.worksheets[:20]:
                chunks.append(f"[工作表: {sheet.title}]")
                for row in sheet.iter_rows(max_row=500, values_only=True):
                    chunks.append("\t".join("" if value is None else str(value) for value in row[:50]))
                    if sum(map(len, chunks)) >= max_chars:
                        break
                if sum(map(len, chunks)) >= max_chars:
                    break
            workbook.close()
            return "\n".join(chunks)[:max_chars]
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
                rows = csv.reader(handle)
                return "\n".join("\t".join(row[:50]) for _, row in zip(range(500), rows))[:max_chars]
    except Exception as exc:
        return f"[文本提取失败: {type(exc).__name__}]"
    return ""


def image_metadata(path: Path) -> str:
    if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp", ".gif", ".tif", ".tiff"}:
        return ""
    try:
        with Image.open(path) as image:
            return f"图像尺寸 {image.width}x{image.height}，格式 {image.format or path.suffix[1:]}"
    except Exception:
        return ""

